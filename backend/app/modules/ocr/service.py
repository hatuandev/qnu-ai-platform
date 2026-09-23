"""OCR Service with Engine Routing, Resilient Fallback, and Studio Intelligence."""

from __future__ import annotations

import hashlib
import io
import os
import re
import time
from pathlib import Path
from typing import Any

import openpyxl
import structlog
from PIL import Image, ImageDraw
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException
from app.core.storage import storage_service
from app.modules.ocr.adapters.base import BaseOCRAdapter
from app.modules.ocr.adapters.docling_adapter import DoclingOCRAdapter
from app.modules.ocr.adapters.easyocr_adapter import EasyOCRAdapter
from app.modules.ocr.adapters.gemini_adapter import GeminiOCRAdapter
from app.modules.ocr.adapters.mistral_adapter import MistralOCRAdapter
from app.modules.ocr.adapters.mock_adapter import MockOCRAdapter
from app.modules.ocr.adapters.pymupdf_adapter import PyMuPDFOCRAdapter
from app.modules.ocr.layout_detector import smart_layout_detector
from app.modules.ocr.models import OCRJobLog
from app.modules.ocr.schemas import (
    OCREngineResponse,
    OCRExtractResponse,
    OCRPageResult,
    StudioOCRPageResponse,
    StudioOCRParseResponse,
    StudioOCRRegion,
)

logger = structlog.get_logger(__name__)

#: Install hints shown when an optional engine is requested but missing.
_ENGINE_INSTALL_HINTS: dict[str, str] = {
    "docling": "uv add docling",
    "easyocr": "uv add easyocr",
}


class OCRService:
    """Service orchestrating OCR document recognition with automated fallback."""

    def __init__(self) -> None:
        self._adapters: dict[str, BaseOCRAdapter] = {
            "pymupdf_ocr": PyMuPDFOCRAdapter(),
            "gemini_ocr": GeminiOCRAdapter(),
            "mistral_ocr": MistralOCRAdapter(),
            "docling": DoclingOCRAdapter(),
            "easyocr": EasyOCRAdapter(),
            "mock_ocr": MockOCRAdapter(),
        }
        self._adapters["gemini"] = self._adapters["gemini_ocr"]
        self._adapters["gemini_vision"] = self._adapters["gemini_ocr"]
        self._adapters["mistral"] = self._adapters["mistral_ocr"]
        self.default_engine = "pymupdf_ocr"

    def list_engines(self) -> list[OCREngineResponse]:
        """List registered OCR engines with real availability status."""
        catalog = [
            ("eng_pymupdf", "pymupdf_ocr", "pymupdf", ["pdf", "png", "jpg", "tables"], 0.95, True),
            ("eng_gemini", "gemini_ocr", "gemini", ["pdf", "png", "jpg", "scan", "tables", "vietnamese", "multimodal"], 0.99, False),
            ("eng_mistral", "mistral_ocr", "mistral", ["pdf", "png", "jpg", "scan", "tables", "vietnamese"], 0.98, False),
            ("eng_docling", "docling", "docling", ["docx", "xlsx", "pdf", "tables", "layout", "markdown"], 0.97, False),
            ("eng_easyocr", "easyocr", "easyocr", ["png", "jpg", "scan", "stamp"], 0.90, False),
            ("eng_mock", "mock_ocr", "mock", ["pdf", "text", "fallback"], 0.98, False),
        ]
        engines: list[OCREngineResponse] = []
        for engine_id, name, engine_type, capabilities, confidence, is_default in catalog:
            adapter = self._adapters[name]
            engines.append(
                OCREngineResponse(
                    id=engine_id,
                    name=name,
                    display_name=adapter.display_name,
                    engine_type=engine_type,
                    provider_category="cloud" if name in ("gemini_ocr", "mistral_ocr") else "local",
                    capabilities=capabilities,
                    avg_confidence=confidence,
                    is_active=adapter.is_available(),
                    is_default=is_default,
                )
            )
        return engines

    def _resolve_adapter(
        self, name_or_model: str, default_gemini_model: str | None = None, **kwargs: Any
    ) -> BaseOCRAdapter:
        """Resolve or dynamically instantiate the appropriate OCR adapter for a model name or engine name."""
        norm = (name_or_model or "").strip().lower()

        # If gemini_ocr/gemini/gemini_vision is requested, use dynamic model if configured
        if norm in ("gemini_ocr", "gemini", "gemini_vision"):
            if default_gemini_model and (
                "gemini" in default_gemini_model.lower()
                or default_gemini_model.lower().startswith("gemma")
            ):
                return GeminiOCRAdapter(model_name=default_gemini_model)
            return self._adapters["gemini_ocr"]

        # Check existing registered adapter keys
        if norm in self._adapters:
            return self._adapters[norm]

        # Check for Google Gemini models
        if "gemini" in norm or norm.startswith("gemma"):
            return GeminiOCRAdapter(model_name=name_or_model)

        # Check for Mistral OCR
        if "mistral" in norm:
            if norm != "mistral_ocr":
                return MistralOCRAdapter(model_name=name_or_model)
            return self._adapters["mistral_ocr"]

        # Check for Docling
        if "docling" in norm:
            return self._adapters["docling"]

        # Check for EasyOCR
        if "easyocr" in norm:
            return self._adapters["easyocr"]

        # Default fallback adapter
        return self._adapters.get(norm) or self._adapters[self.default_engine]

    async def _get_system_defaults(self, session: AsyncSession | None) -> dict[str, Any]:
        """Retrieve system_model_defaults dictionary from database."""
        if session:
            try:
                from sqlalchemy import select

                from app.modules.modelops.models import ModelProviderConfig

                stmt = select(ModelProviderConfig).where(ModelProviderConfig.id == "system_model_defaults")
                res = await session.execute(stmt)
                cfg = res.scalar_one_or_none()
                if cfg and cfg.extra_config and "defaults" in cfg.extra_config:
                    return cfg.extra_config["defaults"]
            except Exception as exc:
                logger.debug("Could not fetch system_model_defaults: %s", exc)
        return {}

    async def _get_active_combo_chain(self, session: AsyncSession | None) -> list[dict[str, Any]]:
        """Retrieve configured OCR combo chain from PostgreSQL system_model_defaults."""
        defaults = await self._get_system_defaults(session)
        chain = defaults.get("ocr_combo_chain")
        if chain and isinstance(chain, list) and len(chain) > 0:
            return chain

        from app.modules.modelops.services.model_catalog_service import DEFAULT_QNU_OCR_COMBO_CHAIN

        return DEFAULT_QNU_OCR_COMBO_CHAIN

    async def extract_document(
        self,
        session: AsyncSession,
        content: bytes,
        filename: str,
        engine_name: str | None = None,
        tenant_id: str = "tenant_qnu",
    ) -> OCRExtractResponse:
        """Extract text from document bytes with automatic graceful fallback or multi-model combo failover."""
        start_time = time.perf_counter()
        requested = (engine_name or "").strip().lower() or "auto"

        defaults = await self._get_system_defaults(session)
        default_ocr_mode = defaults.get("default_ocr_mode")
        default_ocr_model = defaults.get("default_ocr_model")

        if requested in ("auto", "none", ""):
            # If system defaults is configured in combo mode, route automatically to combo chain
            if default_ocr_mode == "combo":
                return await self._extract_combo_chain(session, content, filename, tenant_id, start_time)
            return await self._extract_auto(
                session, content, filename, tenant_id, start_time, default_ocr_model=default_ocr_model
            )

        if requested == "combo":
            return await self._extract_combo_chain(session, content, filename, tenant_id, start_time)

        try:
            adapter = self._resolve_adapter(requested, default_gemini_model=default_ocr_model)
        except TypeError:
            adapter = self._resolve_adapter(requested)

        fallback_triggered = False
        fallback_engine: str | None = None
        cascade_trace: list[str] = []

        if not adapter.is_available():
            if requested in ("gemini_ocr", "gemini", "gemini_vision", "mistral_ocr", "mistral"):
                logger.warning("cloud_ocr_missing_key_fallback_to_local", requested=requested)
                for local_cand in ("easyocr", "docling", self.default_engine):
                    cand_adapter = self._adapters.get(local_cand)
                    if cand_adapter and cand_adapter.is_available():
                        adapter = cand_adapter
                        fallback_triggered = True
                        fallback_engine = local_cand
                        cascade_trace.append(f"{requested}: Thiếu API Key -> Chuyển sang {local_cand}")
                        break
            else:
                hint = _ENGINE_INSTALL_HINTS.get(requested, "lien he quan tri vien")
                raise AppException(
                    f"Engine OCR '{requested}' chua san sang (thieu dependency). Cai dat voi: {hint}",
                    code="ocr_engine_unavailable",
                    status_code=400,
                    details={"engine": requested, "install_hint": hint},
                )

        result_dict = None
        error_msg = None

        try:
            result_dict = await adapter.extract(content, filename)
        except Exception as exc:
            logger.warning("primary_ocr_failed_triggering_fallback", engine=adapter.name, error=str(exc))
            fallback_triggered = True

            # Try combo chain first
            combo_chain = await self._get_active_combo_chain(session)
            for step in combo_chain:
                step_model = step.get("model_name") if isinstance(step, dict) else getattr(step, "model_name", str(step))
                if not step_model or step_model.lower() == requested or step_model.lower() == adapter.name.lower():
                    continue

                fb_adapter = self._resolve_adapter(step_model)
                if not fb_adapter.is_available():
                    continue

                try:
                    fb_res = await fb_adapter.extract(content, filename)
                    raw_text = (fb_res.get("raw_text") or "").strip()
                    if raw_text:
                        result_dict = fb_res
                        result_dict["engine_used"] = step_model
                        fallback_engine = step_model
                        cascade_trace.append(f"{step_model}: Thành công")
                        break
                except Exception as fb_exc:
                    logger.debug("combo_candidate_failed", model=step_model, error=str(fb_exc))
                    continue

            # Fallback to mock safety net if all failed
            if result_dict is None or not (result_dict.get("raw_text") or "").strip():
                fallback_adapter = self._adapters["mock_ocr"]
                try:
                    result_dict = await fallback_adapter.extract(content, filename)
                    result_dict["engine_used"] = fallback_adapter.name
                    fallback_engine = fallback_adapter.name
                    cascade_trace.append("mock_ocr: Kích hoạt lưới an toàn")
                except Exception as fb_exc:
                    error_msg = f"OCR and fallback both failed: {fb_exc}"
                    logger.error("ocr_fallback_failed", error=error_msg)
                    raise AppException(error_msg, status_code=500) from fb_exc

        return await self._build_response(
            session=session,
            result_dict=result_dict or {"engine_used": adapter.name, "raw_text": "", "pages": []},
            filename=filename,
            tenant_id=tenant_id,
            start_time=start_time,
            fallback_triggered=fallback_triggered,
            fallback_engine=fallback_engine,
            cascade_trace=cascade_trace,
            error_msg=error_msg,
        )

    async def _extract_auto(
        self,
        session: AsyncSession,
        content: bytes,
        filename: str,
        tenant_id: str,
        start_time: float,
        default_ocr_model: str | None = None,
    ) -> OCRExtractResponse:
        """Smart routing:
        1. For digital text PDFs: Fast PyMuPDF extraction first (10-30ms).
        2. For scanned PDFs & Images: Prioritize Google Gemini Vision OCR / Mistral OCR (Cloud API, 1-2s).
           If Cloud API keys are not configured, gracefully fall back to Local OCR (EasyOCR/Docling).
        3. For Office files (.docx, .doc, .xlsx, .xls): Prioritize Docling TableFormer.
        """
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "pdf"

        # Special priority order by file format
        if ext in ("docx", "doc", "xlsx", "xls"):
            candidate_order = ("docling", "easyocr", "pymupdf_ocr")
        elif ext in ("png", "jpg", "jpeg", "webp", "bmp", "tiff"):
            candidate_order = ("gemini_ocr", "mistral_ocr", "easyocr", "docling")
        else:
            candidate_order = ("gemini_ocr", "mistral_ocr", "easyocr", "docling")

        # For PDF and text files, run Firecrawl Rust-based PDF Inspector for fast classification
        default_adapter = self._adapters[self.default_engine]
        result_dict = None

        if ext == "pdf":
            from app.modules.knowledge.parsers.pdf_inspector import PDFInspector

            inspection = PDFInspector.inspect_bytes(content)

            # If Firecrawl confirms this is clean digital text (text_based), extract fast without OCR!
            if inspection.can_skip_ocr and inspection.markdown:
                logger.info(
                    "auto_ocr_pdf_inspector_fast_path",
                    pdf_type=inspection.pdf_type,
                    confidence=inspection.confidence,
                    elapsed_ms=round(inspection.processing_time_ms, 2),
                )
                try:
                    result_dict = await default_adapter.extract(content, filename)
                except Exception as exc:
                    logger.warning("default_adapter_extract_failed: %s", exc)
                    result_dict = None

                if not result_dict or not (result_dict.get("raw_text") or "").strip():
                    page_list = [
                        {
                            "page_number": p.page_number,
                            "text": p.markdown,
                            "confidence": inspection.confidence,
                            "blocks": [],
                        }
                        for p in inspection.pages_detail
                    ] or [
                        {
                            "page_number": 1,
                            "text": inspection.markdown,
                            "confidence": inspection.confidence,
                            "blocks": [],
                        }
                    ]
                    result_dict = {
                        "raw_text": inspection.markdown,
                        "pages": page_list,
                        "engine_used": "pdf_inspector_fast_path",
                        "total_pages": inspection.page_count or 1,
                        "latency_ms": inspection.processing_time_ms,
                    }

                return await self._build_response(
                    session=session,
                    result_dict=result_dict,
                    filename=filename,
                    tenant_id=tenant_id,
                    start_time=start_time,
                    fallback_triggered=False,
                    fallback_engine=None,
                    cascade_trace=["pdf_inspector_fast_path: Văn bản số nguyên bản, không cần OCR"],
                    error_msg=None,
                )

        elif ext not in ("png", "jpg", "jpeg", "webp", "bmp", "tiff", "xlsx", "xls"):
            try:
                result_dict = await default_adapter.extract(content, filename)
            except Exception as exc:
                logger.error("auto_ocr_primary_failed", engine=default_adapter.name, error=str(exc))

        # If primary extracted substantial digital text (>80 chars), return immediately!
        if (
            result_dict
            and (result_dict.get("raw_text") or "").strip()
            and len(result_dict.get("raw_text", "").strip()) > 80
        ):
            return await self._build_response(
                session=session,
                result_dict=result_dict,
                filename=filename,
                tenant_id=tenant_id,
                start_time=start_time,
                fallback_triggered=False,
                fallback_engine=None,
                cascade_trace=["digital_text_primary: Trích xuất trực tiếp thành công"],
                error_msg=None,
            )

        # Scanned PDF or images or office files:
        # Loop through candidates in priority order:
        # Dynamic cloud models (Gemini / Mistral) -> Local OCR (EasyOCR / Docling)
        gemini_adapter = (
            GeminiOCRAdapter(model_name=default_ocr_model)
            if default_ocr_model
            and (
                "gemini" in default_ocr_model.lower()
                or default_ocr_model.lower().startswith("gemma")
            )
            else self._adapters["gemini_ocr"]
        )
        mistral_adapter = (
            MistralOCRAdapter(model_name=default_ocr_model)
            if default_ocr_model and "mistral" in default_ocr_model.lower()
            else self._adapters["mistral_ocr"]
        )

        effective_candidates = list(candidate_order)
        if default_ocr_model and "mistral" in default_ocr_model.lower() and "mistral_ocr" in effective_candidates:
            effective_candidates.remove("mistral_ocr")
            effective_candidates.insert(0, "mistral_ocr")

        for candidate in effective_candidates:
            if candidate == "gemini_ocr":
                upgrade = gemini_adapter
            elif candidate == "mistral_ocr":
                upgrade = mistral_adapter
            else:
                upgrade = self._adapters.get(candidate)

            if not upgrade or not upgrade.is_available():
                logger.debug("auto_ocr_candidate_skipped_unavailable", candidate=candidate)
                continue
            try:
                upgraded = await upgrade.extract(content, filename)
            except Exception as exc:
                logger.warning("auto_ocr_candidate_failed_will_try_next", candidate=candidate, error=str(exc))
                continue
            if (upgraded.get("raw_text") or "").strip():
                logger.info(
                    "auto_ocr_upgraded",
                    from_engine=default_adapter.name,
                    to_engine=candidate,
                    filename=filename,
                )
                return await self._build_response(
                    session=session,
                    result_dict=upgraded,
                    filename=filename,
                    tenant_id=tenant_id,
                    start_time=start_time,
                    fallback_triggered=True,
                    fallback_engine=candidate,
                    cascade_trace=[f"{candidate}: Nâng cấp bóc tách thành công"],
                    error_msg=None,
                )

        # If all candidates failed or returned empty, return whatever primary had or empty
        return await self._build_response(
            session=session,
            result_dict=result_dict or {
                "engine_used": default_adapter.name,
                "raw_text": "",
                "pages": [],
                "total_pages": 0,
                "overall_confidence": 0.0,
            },
            filename=filename,
            tenant_id=tenant_id,
            start_time=start_time,
            fallback_triggered=False,
            fallback_engine=None,
            cascade_trace=["auto_scan: Không có văn bản nào được nhận diện"],
            error_msg=None,
        )

    async def _extract_combo_chain(
        self,
        session: AsyncSession,
        content: bytes,
        filename: str,
        tenant_id: str,
        start_time: float,
    ) -> OCRExtractResponse:
        """Sequential Failover across all models configured in ocr_combo_chain."""
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "pdf"

        # Check for clean digital text PDF fast path first
        if ext == "pdf":
            from app.modules.knowledge.parsers.pdf_inspector import PDFInspector

            inspection = PDFInspector.inspect_bytes(content)
            if inspection.can_skip_ocr and inspection.markdown:
                default_adapter = self._adapters[self.default_engine]
                try:
                    result_dict = await default_adapter.extract(content, filename)
                except Exception:
                    result_dict = None

                if not result_dict or not (result_dict.get("raw_text") or "").strip():
                    page_list = [
                        {
                            "page_number": p.page_number,
                            "text": p.markdown,
                            "confidence": inspection.confidence,
                            "blocks": [],
                        }
                        for p in inspection.pages_detail
                    ] or [
                        {
                            "page_number": 1,
                            "text": inspection.markdown,
                            "confidence": inspection.confidence,
                            "blocks": [],
                        }
                    ]
                    result_dict = {
                        "raw_text": inspection.markdown,
                        "pages": page_list,
                        "engine_used": "pdf_inspector_fast_path",
                        "total_pages": inspection.page_count or 1,
                        "latency_ms": inspection.processing_time_ms,
                    }

                return await self._build_response(
                    session=session,
                    result_dict=result_dict,
                    filename=filename,
                    tenant_id=tenant_id,
                    start_time=start_time,
                    fallback_triggered=False,
                    fallback_engine=None,
                    cascade_trace=["pdf_inspector_fast_path: Văn bản số nguyên bản, không cần OCR"],
                    error_msg=None,
                )

        # Scanned PDF, image, or document requiring OCR:
        combo_chain = await self._get_active_combo_chain(session)
        cascade_trace: list[str] = []
        result_dict = None
        fallback_triggered = False
        fallback_engine = None

        for idx, step in enumerate(combo_chain):
            step_model = step.get("model_name") if isinstance(step, dict) else getattr(step, "model_name", str(step))
            if not step_model:
                continue

            adapter = self._resolve_adapter(step_model)
            if not adapter.is_available():
                msg = f"{step_model}: Bỏ qua (Chưa có API Key hoặc thiếu gói phụ thuộc)"
                cascade_trace.append(msg)
                continue

            try:
                step_start = time.perf_counter()
                res = await adapter.extract(content, filename)
                raw_text = (res.get("raw_text") or "").strip()
                if raw_text:
                    step_ms = round((time.perf_counter() - step_start) * 1000, 1)
                    msg = f"{step_model}: Thành công ({step_ms}ms, {len(raw_text)} ký tự)"
                    cascade_trace.append(msg)
                    result_dict = res
                    result_dict["engine_used"] = step_model
                    if idx > 0:
                        fallback_triggered = True
                        fallback_engine = step_model
                    break
                else:
                    cascade_trace.append(f"{step_model}: Kết quả bóc tách trống")
            except Exception as exc:
                exc_str = str(exc)
                is_quota = any(kw in exc_str.lower() for kw in ("429", "quota", "rate limit", "resource exhausted"))
                reason = "Hết hạn ngạch (429 Quota Exceeded)" if is_quota else f"Lỗi: {exc_str[:80]}"
                cascade_trace.append(f"{step_model}: Thất bại ({reason}) -> Tự động chuyển mô hình tiếp theo")
                logger.warning(
                    "combo_ocr_step_failed_continuing_sequence",
                    step=idx + 1,
                    model=step_model,
                    error=exc_str,
                    is_quota=is_quota,
                )
                continue

        # If all items in combo_chain failed, try local fallback
        if result_dict is None or not (result_dict.get("raw_text") or "").strip():
            for local_cand in ("easyocr", "docling", self.default_engine):
                cand_adapter = self._adapters[local_cand]
                if cand_adapter.is_available():
                    try:
                        res = await cand_adapter.extract(content, filename)
                        if (res.get("raw_text") or "").strip():
                            cascade_trace.append(f"{local_cand}: Cứu nguy nội bộ thành công")
                            result_dict = res
                            result_dict["engine_used"] = local_cand
                            fallback_triggered = True
                            fallback_engine = local_cand
                            break
                    except Exception as rescue_exc:
                        logger.debug("local_rescue_failed", candidate=local_cand, error=str(rescue_exc))
                        continue

        if result_dict is None:
            result_dict = {
                "engine_used": "none",
                "raw_text": "",
                "pages": [],
                "total_pages": 0,
                "overall_confidence": 0.0,
            }

        return await self._build_response(
            session=session,
            result_dict=result_dict,
            filename=filename,
            tenant_id=tenant_id,
            start_time=start_time,
            fallback_triggered=fallback_triggered,
            fallback_engine=fallback_engine,
            cascade_trace=cascade_trace,
            error_msg=None,
        )

    async def _build_response(
        self,
        session: AsyncSession,
        result_dict: dict[str, Any],
        filename: str,
        tenant_id: str,
        start_time: float,
        fallback_triggered: bool,
        fallback_engine: str | None = None,
        cascade_trace: list[str] | None = None,
        error_msg: str | None = None,
    ) -> OCRExtractResponse:
        """Persist the audit log and shape the final OCR response."""
        latency_ms = round((time.perf_counter() - start_time) * 1000, 2)
        total_pages = int(result_dict.get("total_pages", 1))
        confidence = float(result_dict.get("overall_confidence", 0.90))

        job_log = OCRJobLog(
            tenant_id=tenant_id,
            filename=filename,
            engine_used=result_dict.get("engine_used", self.default_engine),
            total_pages=total_pages,
            confidence=confidence,
            latency_ms=latency_ms,
            status="success",
            error_message=error_msg,
        )
        try:
            session.add(job_log)
            await session.commit()
        except Exception as db_exc:
            logger.warning("failed_to_save_ocr_job_log", error=str(db_exc))
            await session.rollback()

        pages = [
            OCRPageResult(
                page_number=p["page_number"],
                extracted_text=p["extracted_text"],
                confidence=p["confidence"],
                word_count=p["word_count"],
                line_count=p["line_count"],
                has_tables=p.get("has_tables", False),
                blocks=p.get("blocks", []) or [],
            )
            for p in result_dict.get("pages", [])
        ]

        return OCRExtractResponse(
            filename=filename,
            success=True,
            engine_used=result_dict.get("engine_used", self.default_engine),
            fallback_triggered=fallback_triggered,
            fallback_engine=fallback_engine,
            cascade_trace=cascade_trace or [],
            total_pages=total_pages,
            overall_confidence=confidence,
            latency_ms=latency_ms,
            raw_text=result_dict.get("raw_text", ""),
            pages=pages,
        )

    # --- Studio OCR Interactive Methods ---

    async def parse_studio_document(
        self,
        content: bytes,
        filename: str,
        engine_name: str | None = None,
        tenant_id: str = "tenant_qnu",
        workspace_id: str = "workspace_qnu",
    ) -> StudioOCRParseResponse:
        """Execute OCR recognition on all pages and extract layout bounding boxes for Studio UI."""
        start_time = time.perf_counter()
        file_hash = hashlib.md5(content).hexdigest()[:10]
        ext = os.path.splitext(filename)[1].lower().lstrip(".")

        requested = (engine_name or "").strip().lower() or "auto"
        adapter = self._adapters.get(requested)
        if adapter is None or not adapter.is_available():
            adapter = self._adapters[self.default_engine]

        try:
            ocr_res = await adapter.extract(content, filename)
        except Exception as exc:
            logger.warning("studio_ocr_adapter_failed_fallback", error=str(exc))
            adapter = self._adapters[self.default_engine]
            ocr_res = await adapter.extract(content, filename)

        pages_dict = {p["page_number"]: p for p in ocr_res.get("pages", [])}
        rendered_pages: list[StudioOCRPageResponse] = []

        if ext in ["pdf", "png", "jpg", "jpeg", "webp", "bmp"]:
            try:
                import pymupdf as fitz
                doc = fitz.open(stream=content, filetype=ext if ext != "pdf" else "pdf")
                for i, page in enumerate(doc):
                    pnum = i + 1
                    img_filename = f"page_{pnum}.jpg"
                    cache_key = f"ocr_cache/{file_hash}/{img_filename}"

                    img_bytes = await storage_service.get(cache_key)
                    if not img_bytes:
                        pix = page.get_pixmap(dpi=120)
                        img_bytes = pix.tobytes("jpeg")
                        await storage_service.save(cache_key, img_bytes)

                    p_info = pages_dict.get(pnum)
                    page_md = p_info["extracted_text"] if p_info else ""
                    lines = [ln.strip() for ln in page_md.splitlines() if ln.strip() and not ln.strip().startswith("|")]
                    title = lines[0] if lines else f"Trang {pnum}"

                    cv_regions = smart_layout_detector.detect_layout_regions(
                        img_bytes, page_md, page_number=pnum, fitz_page=page
                    )
                    filtered_regions = [r for r in cv_regions if r.get("type") != "image"]
                    has_table = any(r["type"] == "table" for r in filtered_regions) or ("|" in page_md)
                    is_signed = any(r["type"] == "signature" for r in filtered_regions)

                    regions_models = [
                        StudioOCRRegion(
                            type=r["type"],
                            label=r.get("label", r["type"]),
                            text=r.get("text", ""),
                            top=float(r["top"]),
                            left=float(r["left"]),
                            width=float(r["width"]),
                            height=float(r["height"]),
                        )
                        for r in filtered_regions
                    ]

                    img_url = f"/api/v1/ocr/studio/page-image/{file_hash}/{img_filename}"
                    rendered_pages.append(
                        StudioOCRPageResponse(
                            page_number=pnum,
                            title=title[:90],
                            is_signed=is_signed,
                            has_table=has_table,
                            image_url=img_url,
                            markdown=page_md,
                            raw_text=re.sub(r"[*_#`]", "", page_md).strip(),
                            regions=regions_models,
                            dimensions={"width": int(page.rect.width), "height": int(page.rect.height)},
                            word_count=len(page_md.split()),
                            line_count=len(page_md.splitlines()),
                        )
                    )
            except Exception as pdf_err:
                logger.error("error_rendering_pdf_pages_in_studio", error=str(pdf_err))

        if not rendered_pages:
            total_pages = max(1, ocr_res.get("total_pages", 1))
            for pnum in range(1, total_pages + 1):
                img_filename = f"page_{pnum}.jpg"
                cache_key = f"ocr_cache/{file_hash}/{img_filename}"
                img_bytes = await storage_service.get(cache_key)
                if not img_bytes:
                    img_bytes = self._generate_card_preview_image(filename, img_filename)
                    await storage_service.save(cache_key, img_bytes)

                p_info = pages_dict.get(pnum)
                page_md = p_info["extracted_text"] if p_info else ocr_res.get("raw_text", "")
                lines = [ln.strip() for ln in page_md.splitlines() if ln.strip() and not ln.strip().startswith("|")]
                title = lines[0] if lines else f"Trang {pnum}"

                img_url = f"/api/v1/ocr/studio/page-image/{file_hash}/{img_filename}"
                rendered_pages.append(
                    StudioOCRPageResponse(
                        page_number=pnum,
                        title=title[:90],
                        is_signed=False,
                        has_table="|" in page_md,
                        image_url=img_url,
                        markdown=page_md,
                        raw_text=page_md,
                        regions=[],
                        dimensions={"width": 1240, "height": 1754},
                        word_count=len(page_md.split()),
                        line_count=len(page_md.splitlines()),
                    )
                )

        sheets_data = None
        if ext in ["xlsx", "xls"]:
            sheets_data = self._extract_xlsx_sheets(content)
            if sheets_data and rendered_pages:
                for idx, sheet in enumerate(sheets_data):
                    if idx < len(rendered_pages):
                        rendered_pages[idx].sheet_data = sheet

        latency_ms = round((time.perf_counter() - start_time) * 1000, 2)
        size_str = (
            f"{len(content) / (1024 * 1024):.1f} MB"
            if len(content) >= 1024 * 1024
            else f"{len(content) / 1024:.1f} KB"
        )

        return StudioOCRParseResponse(
            filename=filename,
            total_pages=len(rendered_pages),
            size=size_str,
            provider="local",
            model=adapter.name,
            latency_ms=latency_ms,
            pages=rendered_pages,
            sheets_data=sheets_data,
        )

    def _generate_card_preview_image(self, filename: str, img_filename: str) -> bytes:
        """Generate a clean visual card placeholder for non-rendered Office / Spreadsheet pages."""
        img = Image.new("RGB", (1240, 1754), color=(248, 250, 252))
        draw = ImageDraw.Draw(img)
        draw.rectangle([60, 60, 1180, 1694], outline=(203, 213, 225), width=3)
        draw.rectangle([60, 60, 1180, 160], fill=(16, 185, 129))
        draw.text((100, 95), "QNU.AI STUDIO -- BAN TRINH CHIEU TAI LIEU", fill=(255, 255, 255))
        draw.text((100, 260), f"Tai lieu: {filename}", fill=(15, 23, 42))
        draw.text((100, 320), f"Dang hien thi: {img_filename}", fill=(100, 116, 139))
        draw.text((100, 420), "Da boc tach toan bo cau truc noi dung.", fill=(5, 150, 105))
        draw.text(
            (100, 470),
            "Vui long xem noi dung Markdown & Bieu mau chi tiet o khung ben phai.",
            fill=(71, 85, 105),
        )

        out = io.BytesIO()
        img.save(out, format="JPEG", quality=90)
        return out.getvalue()

    def _extract_xlsx_sheets(self, content: bytes) -> list[dict[str, Any]]:
        """Extract table sheets from an Excel file bytes."""
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheets_data = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_str = ["" if v is None else str(v) for v in row]
                    if any(cell != "" for cell in row_str):
                        rows.append(row_str)
                if rows:
                    max_cols = max(len(r) for r in rows)
                    padded = [r + [""] * (max_cols - len(r)) for r in rows]
                    sheets_data.append({
                        "name": sheet_name,
                        "rows": padded,
                        "total_rows": len(padded),
                        "total_cols": max_cols,
                        "totalRows": len(padded),
                        "totalCols": max_cols,
                    })
            return sheets_data
        except Exception as exc:
            logger.warning("failed_to_extract_xlsx_sheets", error=str(exc))
            return []

    async def get_studio_page_image(self, file_hash: str, img_filename: str) -> bytes:
        """Stream rendered page image from storage cache with multiple fallback locations."""
        cache_key = f"ocr_cache/{file_hash}/{img_filename}"
        img_bytes = await storage_service.get(cache_key)
        if img_bytes:
            return img_bytes

        local_path = Path("frontend/public/ocr-cache") / file_hash / img_filename
        if local_path.is_file():
            return local_path.read_bytes()

        demo_path = Path("frontend/public/ocr-cache/doc_ts_2026") / img_filename
        if demo_path.is_file():
            return demo_path.read_bytes()

        # Generate on-demand fallback card
        card_bytes = self._generate_card_preview_image(file_hash, img_filename)
        await storage_service.save(cache_key, card_bytes)
        return card_bytes

    async def get_sample_document(self) -> StudioOCRParseResponse:
        """Return the pre-configured 14-page official QNU admissions sample document for instant studio testing."""
        pages: list[StudioOCRPageResponse] = []
        for pnum in range(1, 15):
            img_filename = f"page_{pnum}.jpg"
            img_url = f"/api/v1/ocr/studio/page-image/doc_ts_2026/{img_filename}"

            regions: list[StudioOCRRegion] = []
            if pnum == 1:
                regions.append(StudioOCRRegion(type="header", label="header", text="TRƯỜNG ĐẠI HỌC QUY NHƠN", top=4.5, left=7.2, width=38.0, height=5.5))
                regions.append(StudioOCRRegion(type="header", label="header", text="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", top=4.5, left=54.0, width=40.0, height=5.5))
                regions.append(StudioOCRRegion(type="title", label="title", text="THÔNG TIN TUYỂN SINH ĐẠI HỌC CHÍNH QUY NĂM 2026", top=14.0, left=15.0, width=70.0, height=6.0))
            elif pnum in (3, 4, 5, 8, 9, 10):
                regions.append(StudioOCRRegion(type="table", label="table", text="Bảng mã ngành, chỉ tiêu và tổ hợp xét tuyển", top=20.0, left=6.0, width=88.0, height=60.0))
            elif pnum == 14:
                regions.append(StudioOCRRegion(type="signature", label="signature", text="KT. HIỆU TRƯỞNG - PHÓ HIỆU TRƯỞNG (Đã ký & Đóng dấu)", top=72.0, left=58.0, width=35.0, height=18.0))

            sample_md = f"### Trang {pnum}: Thông Tin Tuyển Sinh Đại Học 2026\n\nNội dung văn bản quy chế tuyển sinh trường Đại học Quy Nhơn."
            if pnum in (3, 4, 5, 8, 9, 10):
                sample_md += "\n\n| STT | Mã ngành | Tên ngành đào tạo | Chỉ tiêu | Tổ hợp xét tuyển |\n|:---:|:---:|:---|:---:|:---:|\n| 1 | 7480201 | Công nghệ thông tin | 180 | A00, A01, D01, D07 |\n| 2 | 7480101 | Khoa học máy tính | 120 | A00, A01, D01 |\n| 3 | 7480108 | Kỹ thuật phần mềm | 100 | A00, A01, D01, D07 |"

            pages.append(
                StudioOCRPageResponse(
                    page_number=pnum,
                    title=f"Thông tin tuyển sinh 2026 - Trang {pnum}",
                    is_signed=(pnum == 14),
                    has_table=(pnum in (3, 4, 5, 8, 9, 10)),
                    image_url=img_url,
                    markdown=sample_md,
                    raw_text=sample_md,
                    regions=regions,
                    dimensions={"width": 1240, "height": 1754},
                    word_count=len(sample_md.split()),
                    line_count=len(sample_md.splitlines()),
                )
            )

        return StudioOCRParseResponse(
            filename="Thong tin tuyen sinh dai hoc 2026_Lan2-1 (1).pdf",
            total_pages=14,
            size="3.2 MB",
            provider="local",
            model="pymupdf_ocr",
            latency_ms=45.0,
            pages=pages,
            sheets_data=None,
        )
