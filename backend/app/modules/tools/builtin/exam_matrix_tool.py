"""Exam Matrix Exporter Tool — Generates Excel .xlsx files based on Bloom's Taxonomy."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.core.config import settings
from app.modules.tools.builtin.base import BaseTool


class ExamMatrixExporterTool(BaseTool):
    """Tool that formats exam matrix according to Bloom's Taxonomy into an Excel spreadsheet."""

    @property
    def name(self) -> str:
        return "export_exam_matrix"

    @property
    def display_name(self) -> str:
        return "Xuất Ma trận Đề thi Excel (.xlsx) Chuẩn Bloom"

    @property
    def description(self) -> str:
        return (
            "Tạo bảng ma trận phân phối câu hỏi và trọng số điểm thi theo 4 cấp độ tư duy Bloom "
            "(Nhận biết, Thông hiểu, Vận dụng, Vận dụng cao) chuẩn quy chế khảo thí Trường Đại học Quy Nhơn."
        )

    @property
    def category(self) -> str:
        return "question_bank"

    def get_openapi_schema(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "description": self.description,
            "parameters": {
                "type": "object",
                "properties": {
                    "course_name": {
                        "type": "string",
                        "description": "Tên học phần (ví dụ: Lập trình Cơ sở dữ liệu, Trí tuệ Nhân tạo)",
                    },
                    "course_code": {
                        "type": "string",
                        "description": "Mã học phần (ví dụ: CS204, IT302)",
                        "default": "QNU101",
                    },
                    "exam_duration_minutes": {
                        "type": "integer",
                        "description": "Thời gian làm bài thi (phút)",
                        "default": 60,
                    },
                    "topics": {
                        "type": "array",
                        "description": "Danh sách các chủ đề kiến thức phân bổ theo thang Bloom",
                        "items": {
                            "type": "object",
                            "properties": {
                                "topic_name": {"type": "string", "description": "Tên chủ đề / Chương"},
                                "recognition_count": {"type": "integer", "description": "Số câu Nhận biết (Cấp 1)"},
                                "comprehension_count": {"type": "integer", "description": "Số câu Thông hiểu (Cấp 2)"},
                                "application_count": {"type": "integer", "description": "Số câu Vận dụng (Cấp 3)"},
                                "advanced_application_count": {"type": "integer", "description": "Số câu Vận dụng cao (Cấp 4)"},
                                "total_score": {"type": "number", "description": "Tổng điểm của chủ đề"},
                            },
                            "required": ["topic_name"],
                        },
                    },
                },
                "required": ["course_name", "topics"],
            },
        }

    async def execute(
        self, parameters: dict[str, Any], context: dict[str, Any] | None = None
    ) -> dict[str, Any]:
        course_name = parameters.get("course_name", "Học phần mẫu")
        course_code = parameters.get("course_code", "QNU101")
        duration = parameters.get("exam_duration_minutes", 60)
        topics = parameters.get("topics", [])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ma Trận Đề Thi"

        # Fonts & Styles
        font_header = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
        font_title = Font(name="Times New Roman", size=14, bold=True, color="003366")
        font_sub = Font(name="Times New Roman", size=11, italic=True)
        font_data = Font(name="Times New Roman", size=11)
        font_total = Font(name="Times New Roman", size=11, bold=True)

        fill_qnu_blue = PatternFill(start_color="005696", end_color="005696", fill_type="solid")
        fill_sub_blue = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

        thin = Side(border_style="thin", color="000000")
        border_all = Border(top=thin, left=thin, right=thin, bottom=thin)

        # Title block
        ws.merge_cells("A1:G1")
        ws["A1"] = "TRƯỜNG ĐẠI HỌC QUY NHƠN — PHÒNG KHẢO THÍ & BẢO ĐẢM CHẤT LƯỢNG"
        ws["A1"].font = Font(name="Times New Roman", size=12, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:G2")
        ws["A2"] = f"MA TRẬN ĐỀ THI KẾT THÚC HỌC PHẦN: {course_name.upper()} ({course_code})"
        ws["A2"].font = font_title
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A3:G3")
        ws["A3"] = f"Thời gian làm bài: {duration} phút | Hình thức thi: Trắc nghiệm khách quan kết hợp Tự luận"
        ws["A3"].font = font_sub
        ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

        # Table Header Row 5
        headers = [
            ("A5", "STT"),
            ("B5", "Chủ đề / Nội dung kiến thức"),
            ("C5", "Nhận biết\n(Cấp độ 1)"),
            ("D5", "Thông hiểu\n(Cấp độ 2)"),
            ("E5", "Vận dụng\n(Cấp độ 3)"),
            ("F5", "Vận dụng cao\n(Cấp độ 4)"),
            ("G5", "Tổng điểm / Tỷ lệ"),
        ]

        for cell_ref, text in headers:
            cell = ws[cell_ref]
            cell.value = text
            cell.font = font_header
            cell.fill = fill_qnu_blue
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_all

        ws.row_dimensions[5].height = 32

        # Data Rows
        current_row = 6
        tot_c1 = 0
        tot_c2 = 0
        tot_c3 = 0
        tot_c4 = 0
        tot_score = 0.0

        for idx, t in enumerate(topics, start=1):
            c1 = int(t.get("recognition_count", 0))
            c2 = int(t.get("comprehension_count", 0))
            c3 = int(t.get("application_count", 0))
            c4 = int(t.get("advanced_application_count", 0))
            sc = float(t.get("total_score", 0.0))

            tot_c1 += c1
            tot_c2 += c2
            tot_c3 += c3
            tot_c4 += c4
            tot_score += sc

            row_data = [
                (f"A{current_row}", idx, "center"),
                (f"B{current_row}", t.get("topic_name", f"Chủ đề {idx}"), "left"),
                (f"C{current_row}", c1, "center"),
                (f"D{current_row}", c2, "center"),
                (f"E{current_row}", c3, "center"),
                (f"F{current_row}", c4, "center"),
                (f"G{current_row}", f"{sc:.1f} đ", "center"),
            ]

            for ref, val, align in row_data:
                cell = ws[ref]
                cell.value = val
                cell.font = font_data
                cell.border = border_all
                cell.alignment = Alignment(horizontal=align, vertical="center")

            current_row += 1

        # Summary Row
        ws[f"A{current_row}"] = ""
        ws[f"B{current_row}"] = "TỔNG CỘNG"
        ws[f"C{current_row}"] = tot_c1
        ws[f"D{current_row}"] = tot_c2
        ws[f"E{current_row}"] = tot_c3
        ws[f"F{current_row}"] = tot_c4
        ws[f"G{current_row}"] = f"{tot_score:.1f} đ"

        for col_letter in ["A", "B", "C", "D", "E", "F", "G"]:
            cell = ws[f"{col_letter}{current_row}"]
            cell.font = font_total
            cell.fill = fill_sub_blue
            cell.border = border_all
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws[f"B{current_row}"].alignment = Alignment(horizontal="left", vertical="center")

        # Column widths
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 18

        # Save to Artifacts directory
        artifacts_dir = Path(settings.LOCAL_STORAGE_PATH) / "artifacts"
        artifacts_dir.mkdir(parents=True, exist_ok=True)
        clean_name = course_code.lower().replace(" ", "_")
        filename = f"ma_tran_de_thi_{clean_name}.xlsx"
        file_path = artifacts_dir / filename
        wb.save(str(file_path))

        return {
            "status": "generated",
            "file_name": filename,
            "file_path": str(file_path),
            "course_name": course_name,
            "course_code": course_code,
            "total_questions": tot_c1 + tot_c2 + tot_c3 + tot_c4,
            "total_score": tot_score,
            "bloom_distribution": {
                "recognition": tot_c1,
                "comprehension": tot_c2,
                "application": tot_c3,
                "advanced_application": tot_c4,
            },
            "size_bytes": file_path.stat().st_size,
        }
