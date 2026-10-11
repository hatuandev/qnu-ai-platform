"""Opt-in live worker/ModelOps/Qdrant verification in an isolated test collection.

No reset, no user collection mutation, no secrets in output, no production mocks.
Creates two clearly labelled fixture documents; leaves the fixture for UI review.
"""

import asyncio
import io
import sys
import uuid

import httpx
from docx import Document

from app.core.config import settings


async def main():
    prefix = settings.API_PREFIX
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8001", timeout=120) as client:
        async def api(method, path, **kwargs):
            response = await client.request(method, f"{prefix}{path}", **kwargs)
            if response.is_error:
                raise RuntimeError(f"{method} {path}: HTTP {response.status_code}: {response.text[:700]}")
            return response.json()

        await api("POST", "/auth/login", json={"access_key": settings.DEV_ACCESS_PASSWORD})
        defaults = await api("GET", "/modelops/defaults")
        model = next((item for item in defaults["available_embeddings"]
                      if item["model_name"] == defaults["defaults"]["default_embedding_model"]
                      and item["provider_id"] == defaults["defaults"]["default_embedding_provider_id"]), None)
        assert model, "Configure an embedding default in ModelOps before running this verification."
        tag = uuid.uuid4().hex[:8]
        collection = await api("GET", f"/knowledge/collections/{sys.argv[1]}") if len(sys.argv) > 1 else await api("POST", "/knowledge/collections", json={
            "name": f"Kiểm thử vector #354 {tag}", "module_code": "admissions",
            "description": "FIXTURE kiểm thử; không liên kết trợ lý và không phải tài liệu chính thức.",
        })
        assert collection["name"].startswith("Kiểm thử vector #354"), "Only isolated #354 test collections are allowed"
        col = collection["id"]
        print(f"Fixture collection: {col}", flush=True)
        existing = await api("GET", f"/knowledge/collections/{col}/bindings")
        doc_ids = [item["repository_document_id"] for item in existing]
        for number in (() if existing else (1, 2)):
            doc = Document()
            doc.add_heading(f"Tài liệu kiểm thử vector {tag} / {number}", 0)
            doc.add_paragraph("Đây là dữ liệu kiểm thử luồng Markdown và vector, không dùng để tư vấn tuyển sinh.")
            doc.add_paragraph(f"Mục {number}. Nội dung độc lập dùng xác minh nạp đúng tài liệu đã chọn.")
            stream = io.BytesIO()
            doc.save(stream)
            uploaded = await api("POST", "/documents/intake", files={"file": (
                f"vector354_{tag}_{number}.docx", stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )}, data={"title": f"FIXTURE vector354 {tag} / {number}"})
            doc_ids.append(uploaded["document_id"])
        if not existing:
            batch = await api("POST", f"/knowledge/collections/{col}/imports", json={
                "document_ids": doc_ids, "parser_profile": "standard", "chunk_strategy": "ClauseBasedChunker",
            })
            for _ in range(90):
                batch = await api("GET", f"/knowledge/collections/{col}/imports/{batch['id']}")
                if batch["status"] in ("completed", "partial_failed", "review_required", "failed"):
                    break
                await asyncio.sleep(1)
            assert batch["succeeded_count"] == 2, batch
        sources = await api("GET", f"/knowledge/collections/{col}/bindings")
        assert len(sources) == 2 and all(item["source_status"] == "ready" for item in sources)
        revisions = {item["id"]: item["active_source_revision_id"] for item in sources}
        generation = await api("POST", f"/knowledge/collections/{col}/vector-generations", json={
            "embedding_model": model["model_name"], "provider_id": model["provider_id"],
            "embedding_dimension": model.get("dimension"),
        })
        gen = generation["id"]
        print(f"Real embedding: {generation['embedding_model']}; generation: {gen}", flush=True)
        for source in sources:
            queued = await api("POST", f"/knowledge/collections/{col}/vector-load", json={
                "generation_id": gen, "items": [{"binding_id": source["id"], "source_revision_id": revisions[source["id"]]}],
            })
            assert len(queued) == 1 and queued[0]["payload"]["binding_id"] == source["id"]
            for _ in range(90):
                state = await api("GET", f"/knowledge/collections/{col}/vector-load")
                job = next(item for item in state["jobs"] if item["id"] == queued[0]["id"])
                if job["status"] in ("completed", "failed", "cancelled"):
                    break
                await asyncio.sleep(1)
            assert job["status"] == "completed", job
            assert job["result"]["parity_passed"] is True
            print(f"Selected-source worker completed: {source['id']}; chunks={job['result']['chunk_count']}", flush=True)
            generations = await api("GET", f"/knowledge/collections/{col}/vector-generations")
            assert not any(item["status"] == "active" for item in generations), "Must not auto-publish"
            if source == sources[0]:
                assert len(state["indexes"]) == 1, "Loading one source must not load the second"
                blocked = await client.post(f"{prefix}/knowledge/collections/{col}/vector-generations/{gen}/activate")
                assert blocked.status_code == 409, "Incomplete collection must not publish"
        final_sources = await api("GET", f"/knowledge/collections/{col}/bindings")
        assert {item["id"]: item["active_source_revision_id"] for item in final_sources} == revisions, "Embedding must not OCR or create another source revision"
        published = await api("POST", f"/knowledge/collections/{col}/vector-generations/{gen}/activate")
        assert published["activated_generation_id"] == gen
        final_sources = await api("GET", f"/knowledge/collections/{col}/bindings")
        assert all(item["active_index_revision_id"] for item in final_sources)
        print(f"PASS: worker -> real embedding -> Qdrant parity -> explicit publish. UI: http://localhost:3000/knowledge/{col}", flush=True)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    asyncio.run(main())
