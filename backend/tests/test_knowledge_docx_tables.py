"""Tests for DocxParser, Table Extraction, and Markdown Table Cleaning."""

from __future__ import annotations

import io

import docx
import pytest

from app.modules.knowledge.cleaner import clean_markdown_text
from app.modules.knowledge.parsers.office_parser import DocxParser
from app.modules.ocr.layout_detector import SmartLayoutDetector


def _create_sample_docx() -> bytes:
    """Tạo tệp DOCX mẫu có tiêu đề hành chính, đoạn văn và bảng biểu tuyển sinh."""
    doc = docx.Document()

    # Bảng hành chính Quốc hiệu / Tiêu ngữ
    hdr_table = doc.add_table(rows=1, cols=2)
    hdr_table.rows[0].cells[0].text = "BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC QUY NHƠN\nSố: 123/TB-ĐHQN"
    hdr_table.rows[0].cells[1].text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\nBình Định, ngày 15 tháng 3 năm 2026"

    # Tiêu đề
    p_title = doc.add_paragraph("THÔNG BÁO TUYỂN SINH ĐẠI HỌC CHÍNH QUY NĂM 2026")
    p_title.style = "Heading 1"

    # Đoạn văn
    doc.add_paragraph("1. Đối tượng dự tuyển: Thí sinh đã tốt nghiệp THPT hoặc tương đương.")
    doc.add_paragraph("2. Tổng chỉ tiêu dự kiến: 4800 chỉ tiêu.")

    # Bảng ngành xét tuyển
    table = doc.add_table(rows=4, cols=4)
    headers = ["STT", "Mã xét tuyển", "Tên ngành đào tạo", "Tổ hợp môn xét tuyển"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    data = [
        ["1", "7140114", "Quản lý giáo dục", "A00, A01, D01\nC00, C04"],
        ["2", "7140201", "Giáo dục Mầm non", "M00, M01"],
        ["3", "7140202", "Giáo dục Tiểu học", "A00, C00, D01"],
    ]
    for r_idx, row_data in enumerate(data, start=1):
        for c_idx, val in enumerate(row_data):
            table.rows[r_idx].cells[c_idx].text = val

    # Phần kết thúc
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.rows[0].cells[0].text = "Nơi nhận:\n- Các khoa;\n- Lưu: VT, ĐT."
    sign_table.rows[0].cells[1].text = "HIỆU TRƯỞNG\nĐã ký\nPGS.TS. Đoàn Đức Tùng"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_docx_parser_sequential_extraction_and_tables() -> None:
    """Kiểm tra DocxParser trích xuất tuần tự văn bản và bảng biểu đầy đủ."""
    docx_bytes = _create_sample_docx()
    parser = DocxParser()

    parsed = await parser.parse(docx_bytes, "tuyen_sinh_2026.docx")

    assert parsed.raw_text
    assert len(parsed.tables) >= 1

    # Kiểm tra bảng hành chính chuyển thành blockquote NĐ 30
    assert "> **BỘ GIÁO DỤC VÀ ĐÀO TẠO" in parsed.raw_text
    assert "Độc lập - Tự do - Hạnh phúc" in parsed.raw_text

    # Kiểm tra đoạn văn được giữ đúng vị trí
    assert "4800 chỉ tiêu" in parsed.raw_text

    # Kiểm tra bảng ngành đào tạo được trích xuất thành Markdown table chuẩn
    assert "| STT | Mã xét tuyển | Tên ngành đào tạo | Tổ hợp môn xét tuyển |" in parsed.raw_text
    assert "| 1 | 7140114 | Quản lý giáo dục |" in parsed.raw_text
    assert "| 2 | 7140201 | Giáo dục Mầm non |" in parsed.raw_text
    assert "| 3 | 7140202 | Giáo dục Tiểu học |" in parsed.raw_text

    # Kiểm tra ô nhiều dòng được xử lý thành <br> hoặc ; không làm vỡ hàng
    assert "\n1 | 7140114" not in parsed.raw_text
    assert "<br>" in parsed.raw_text or "; " in parsed.raw_text

    # Tuyệt đối không có chuỗi giữ chỗ
    assert "Bảng biểu dữ liệu số hóa" not in parsed.raw_text


def test_smart_layout_format_table_markdown() -> None:
    """Kiểm tra hàm _format_table_markdown chuyển đổi ma trận ô thành bảng Markdown."""
    rows = [
        ["STT", "Mã ngành", "Tên ngành", "Chỉ tiêu"],
        ["1", "7480201", "Công nghệ thông tin", "250"],
        ["2", "7480101", "Khoa học máy tính", "100"],
    ]
    md = SmartLayoutDetector._format_table_markdown(rows)

    assert "| STT | Mã ngành | Tên ngành | Chỉ tiêu |" in md
    assert "| :---: | :--- | :--- | :--- |" in md
    assert "| 1 | 7480201 | Công nghệ thông tin | 250 |" in md
    assert "| 2 | 7480101 | Khoa học máy tính | 100 |" in md


def test_clean_markdown_text_table_cleaning() -> None:
    """Kiểm tra clean_markdown_text chuẩn hóa bảng và nối bảng."""
    raw = (
        "Một đoạn văn bản mở đầu.\n\n"
        "| Cột 1 | Cột 2 |\n"
        "| Dữ liệu 1 | Dữ liệu 2 |\n\n"
        "Trang 1 / 14\n\n"
        "| Cột 1 | Cột 2 |\n"
        "| :--- | :--- |\n"
        "| Dữ liệu 3 | Dữ liệu 4 |\n"
    )
    cleaned = clean_markdown_text(raw)

    assert "Trang 1 / 14" not in cleaned
    assert "| Cột 1 | Cột 2 |" in cleaned
    assert "| Dữ liệu 1 | Dữ liệu 2 |" in cleaned
    assert "| Dữ liệu 3 | Dữ liệu 4 |" in cleaned
