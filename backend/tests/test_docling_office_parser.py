"""Tests for DoclingOfficeParser, native fallback routing, and error handling."""

from __future__ import annotations

import asyncio
import io
import unittest

import docx

from app.modules.knowledge.parsers import get_document_parser
from app.modules.knowledge.parsers.docling_office_parser import DoclingOfficeParser
from app.modules.knowledge.parsers.office_parser import DocxParser


def _sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Văn bản thử nghiệm", level=1)
    doc.add_paragraph("Nội dung đoạn văn phục vụ kiểm tra parser.")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Tiêu đề 1"
    tbl.rows[0].cells[1].text = "Tiêu đề 2"
    tbl.rows[1].cells[0].text = "Giá trị 1"
    tbl.rows[1].cells[1].text = "Giá trị 2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDoclingOfficeParser(unittest.TestCase):
    def test_unsupported_extension_raises_value_error(self) -> None:
        parser = DoclingOfficeParser("unsupported_ext")
        with self.assertRaises(ValueError):
            parser._parse_sync(b"dummy", "test.unsupported_ext")

    def test_parser_factory_routes_office_extensions_to_docling_office_parser(self) -> None:
        for ext in ["docx", "doc", "xlsx", "xls", "pptx", "ppt"]:
            parser = get_document_parser(ext)
            self.assertIsInstance(parser, DoclingOfficeParser)
            self.assertEqual(parser.extension, ext)

    def test_docling_office_parser_fallback_on_missing_dependency(self) -> None:
        """When Docling Office backend cannot import (e.g. missing pypdfium2), falls back seamlessly."""
        parser = DoclingOfficeParser("docx", fallback=DocxParser())
        docx_bytes = _sample_docx_bytes()

        parsed = asyncio.run(parser.parse(docx_bytes, "sample.docx"))

        self.assertTrue(parsed.raw_text)
        self.assertIn("Văn bản thử nghiệm", parsed.raw_text)
        self.assertIn("Tiêu đề 1", parsed.raw_text)
        self.assertIn(parsed.metadata.get("parser"), {"docling", "native_fallback"})

    def test_docling_office_parser_raises_without_fallback_on_import_error(self) -> None:
        parser = DoclingOfficeParser("docx", fallback=None)
        with self.assertRaises(RuntimeError):
            asyncio.run(parser.parse(b"corrupt", "sample.docx"))

    def test_extract_markdown_tables_helper(self) -> None:
        markdown = (
            "# Tiêu đề\n\n"
            "Đoạn văn giới thiệu.\n\n"
            "| Cột A | Cột B |\n"
            "| :--- | :--- |\n"
            "| Dữ liệu 1 | Dữ liệu 2 |\n"
            "| Dữ liệu 3 | Dữ liệu 4 |\n\n"
            "Kết thúc văn bản."
        )
        tables = DoclingOfficeParser._extract_markdown_tables(markdown)
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].headers, ["Cột A", "Cột B"])
        self.assertEqual(len(tables[0].rows), 2)
        self.assertEqual(tables[0].rows[0], ["Dữ liệu 1", "Dữ liệu 2"])


if __name__ == "__main__":
    unittest.main()
